# -*- coding: utf-8 -*-
"""PRD 6절 FR-1: CLI 진입점.

사용법:
    python main.py --input data/sample_1parcel.xlsx --output result.xlsx --site-name 신안군마명리
"""
from __future__ import annotations

import argparse
import logging
from pathlib import Path

from openpyxl import load_workbook

from fallback_registry_pdf import load_registry_areas, normalize_jibun
from ledger_builder import build_ledger
from models import FailureCode, ParcelRecord, QueryStatus
from vworld_client import VWorldClient, VWorldError

PROJECT_ROOT = Path(__file__).resolve().parent.parent
PARENT_ENV = PROJECT_ROOT.parent / ".env"  # 나만의 AI Agent/.env (기존 세션 자료)
LOCAL_ENV = PROJECT_ROOT / ".env"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger("main")


def _load_key_from_labeled_env(path: Path, label_substr: str) -> str | None:
    """'브이월드 인증키 : XXXX' 형식의 기존 .env를 읽는다."""
    if not path.exists():
        return None
    for line in path.read_text(encoding="utf-8").splitlines():
        if label_substr in line and ":" in line:
            return line.split(":", 1)[1].strip()
    return None


def get_vworld_key() -> str:
    # 1) 이 프로젝트 전용 .env (KEY=VALUE 표준 형식)
    if LOCAL_ENV.exists():
        from dotenv import dotenv_values

        values = dotenv_values(LOCAL_ENV)
        if values.get("VWORLD_API_KEY"):
            return values["VWORLD_API_KEY"]

    # 2) 상위 나만의 AI Agent/.env (한글 라벨 형식, 기존 세션 자료 재사용)
    key = _load_key_from_labeled_env(PARENT_ENV, "브이월드")
    if key:
        return key

    raise SystemExit(
        "브이월드 API 키를 찾을 수 없습니다. "
        f"{LOCAL_ENV} 에 VWORLD_API_KEY=... 를 추가하거나 "
        f"{PARENT_ENV} 파일을 확인하세요."
    )


def read_input_parcels(input_path: Path) -> list[tuple[str, str]]:
    """엑셀에서 (소재지, 지번) 목록을 읽는다. 헤더: 소재지, 지번."""
    wb = load_workbook(input_path, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        addr_idx = headers.index("소재지")
        jibun_idx = headers.index("지번")
    except ValueError as exc:
        raise SystemExit(f"입력 파일에 '소재지'/'지번' 헤더가 필요합니다. 현재 헤더: {headers}") from exc

    parcels: list[tuple[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        addr = row[addr_idx]
        jibun = row[jibun_idx]
        if addr and jibun:
            parcels.append((str(addr).strip(), str(jibun).strip()))
    return parcels


def process_parcel(
    client: VWorldClient,
    address: str,
    jibun: str,
    registry_areas: dict[str, float],
) -> ParcelRecord:
    record = ParcelRecord(소재지=address, 지번=jibun)

    query_address = f"{address} {jibun}".strip()
    pnu_result = client.search_pnu(query_address)
    if pnu_result is None:
        record.mark_failure(FailureCode.PNU_NOT_FOUND)
    else:
        record.pnu = pnu_result["pnu"]

    # FR-3: 지목/면적 (1차: VWorld ned API)
    if record.pnu:
        try:
            char_data = client.get_land_characteristics(record.pnu)
            if _apply_land_characteristics(record, char_data):
                record.add_source("VWorld API")
            else:
                # PRD FR-3 알려진 제약: 이 엔드포인트는 INCORRECT_KEY 대신
                # resultCode="" + totalCount="0"으로 미승인 상태를 나타낸다
                # (getLandUseAttr/getIndvdLandPriceAttr와 실패 응답 형식이 다름).
                record.mark_failure(FailureCode.API_KEY_UNAUTHORIZED)
        except VWorldError as exc:
            record.mark_failure(exc.code)

        try:
            use_data = client.get_land_use(record.pnu)
            if _apply_land_use(record, use_data):
                record.add_source("VWorld API")
        except VWorldError as exc:
            record.mark_failure(exc.code)

        try:
            price_data = client.get_individual_land_price(record.pnu)
            if _apply_land_price(record, price_data):
                record.add_source("VWorld API")
        except VWorldError as exc:
            record.mark_failure(exc.code)

    # FR-3-Fallback ②: 면적이 여전히 없으면 등기부등본 CSV 오버라이드 적용
    if record.면적_m2 is None:
        area = registry_areas.get(normalize_jibun(jibun))
        if area is not None:
            record.면적_m2 = area
            record.면적_평 = round(area / 3.305785, 4)
            record.add_source("등기부등본")
            # 등기부에서 값을 구했으니 면적 관련 실패사유는 해소된 것으로 간주하지 않고
            # (지목/용도지역은 여전히 비어있을 수 있으므로) 상태만 부분성공으로 조정
            if record.조회상태 == QueryStatus.FAILED:
                record.조회상태 = QueryStatus.PARTIAL

    if record.면적_m2 is None:
        record.mark_failure(FailureCode.NO_MATCHING_PARCEL)

    if not record.실패사유:
        record.조회상태 = QueryStatus.SUCCESS
    elif record.면적_m2 is not None:
        record.조회상태 = QueryStatus.PARTIAL

    return record


def _extract_fields(data: dict) -> list:
    root = next(iter(data.values()), {}) if isinstance(data, dict) else {}
    fields = root.get("field", root.get("fields", []))
    if isinstance(fields, dict):
        fields = [fields]
    return fields or []


def _apply_land_characteristics(record: ParcelRecord, data: dict) -> bool:
    """지목·면적 반영. 실제 데이터를 채웠으면 True (PRD FR-3)."""
    fields = _extract_fields(data)
    if not fields:
        return False
    item = fields[0]
    record.지목 = item.get("lndcgrCodeNm")
    area = item.get("lndpclAr")
    if area:
        try:
            record.면적_m2 = float(area)
            record.면적_평 = round(record.면적_m2 / 3.305785, 4)
        except (TypeError, ValueError):
            pass
    return bool(record.지목 or record.면적_m2)


def _apply_land_use(record: ParcelRecord, data: dict) -> bool:
    fields = _extract_fields(data)
    if not fields:
        return False
    item = fields[0]
    record.용도지역 = item.get("prposArea1Nm")
    record.용도지구 = item.get("prposArea2Nm")
    record.규제사항_원문 = item.get("etcRegstrDivNm")
    return bool(record.용도지역 or record.용도지구 or record.규제사항_원문)


def _apply_land_price(record: ParcelRecord, data: dict) -> bool:
    fields = _extract_fields(data)
    if not fields:
        return False
    item = fields[0]
    price = item.get("pblntfPclnd")
    if price:
        try:
            record.개별공시지가 = int(float(price))
            return True
        except (TypeError, ValueError):
            pass
    return False


def main() -> None:
    parser = argparse.ArgumentParser(description="사업부지 토지조서 자동 작성 (PRD Phase 1 MVP)")
    parser.add_argument("--input", required=True, help="입력 엑셀 (컬럼: 소재지, 지번)")
    parser.add_argument("--output", required=True, help="출력 엑셀 경로")
    parser.add_argument("--site-name", default="", help="조서 제목에 들어갈 부지명")
    parser.add_argument(
        "--registry-csv",
        default=None,
        help="등기부등본에서 미리 추출한 지번,면적_m2 CSV/TSV (FR-3-Fallback ②)",
    )
    parser.add_argument("--no-cache", action="store_true", help="API 응답 캐시 사용 안 함")
    args = parser.parse_args()

    api_key = get_vworld_key()
    client = VWorldClient(api_key=api_key, use_cache=not args.no_cache)

    registry_areas = load_registry_areas(args.registry_csv) if args.registry_csv else {}
    if args.registry_csv:
        logger.info("등기부등본 CSV 로드: %d건", len(registry_areas))

    parcels = read_input_parcels(Path(args.input))
    logger.info("입력 필지 수: %d", len(parcels))

    records = []
    for address, jibun in parcels:
        logger.info("조회 중: %s %s", address, jibun)
        record = process_parcel(client, address, jibun, registry_areas)
        records.append(record)
        logger.info("  -> 상태=%s 사유=%s", record.조회상태.value, record.failure_text())

    column_mapping_path = PROJECT_ROOT / "config" / "column_mapping.json"
    out_path = build_ledger(
        records,
        output_path=args.output,
        column_mapping_path=column_mapping_path if column_mapping_path.exists() else None,
        site_name=args.site_name,
    )

    success = sum(1 for r in records if r.조회상태 == QueryStatus.SUCCESS)
    partial = sum(1 for r in records if r.조회상태 == QueryStatus.PARTIAL)
    failed = sum(1 for r in records if r.조회상태 == QueryStatus.FAILED)
    logger.info("완료: 성공 %d / 부분성공 %d / 실패 %d -> %s", success, partial, failed, out_path)


if __name__ == "__main__":
    main()
