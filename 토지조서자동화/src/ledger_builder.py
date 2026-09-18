# -*- coding: utf-8 -*-
"""PRD FR-8: 조회 결과 -> 엑셀 조서 조립. FR-6: 면적 합산."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import ParcelRecord, QueryStatus

DEFAULT_COLUMNS = [
    ("소재지", "소재지"),
    ("지번", "지번"),
    ("지목", "지목"),
    ("면적_m2", "면적(㎡)"),
    ("면적_평", "면적(평)"),
    ("소유구분", "소유구분"),
    ("용도지역", "용도지역"),
    ("용도지구", "용도지구"),
    ("규제사항_원문", "규제사항"),
    ("개별공시지가", "개별공시지가"),
    ("조회상태", "조회상태"),
    ("데이터출처", "데이터출처"),
    ("조회시각", "조회시각"),
    ("비고", "비고"),
]


def load_column_mapping(config_path: str | Path | None) -> list[tuple[str, str]]:
    """config/column_mapping.json이 있으면 그 헤더명으로 덮어쓴다 (PRD FR-8)."""
    if not config_path:
        return DEFAULT_COLUMNS
    path = Path(config_path)
    if not path.exists():
        return DEFAULT_COLUMNS
    mapping = json.loads(path.read_text(encoding="utf-8"))
    return [(field, mapping.get(field, header)) for field, header in DEFAULT_COLUMNS]


def _row_values(record: ParcelRecord) -> dict[str, object]:
    return {
        "소재지": record.소재지,
        "지번": record.지번,
        "지목": record.지목 or record.failure_text(),
        "면적_m2": record.면적_m2,
        "면적_평": record.면적_평,
        "소유구분": record.소유구분,
        "용도지역": record.용도지역 or (record.failure_text() if not record.면적_m2 else None),
        "용도지구": record.용도지구,
        "규제사항_원문": record.규제사항_원문,
        "개별공시지가": record.개별공시지가,
        "조회상태": record.조회상태.value,
        "데이터출처": record.source_text(),
        "조회시각": record.조회시각.strftime("%Y-%m-%d %H:%M:%S"),
        "비고": record.비고,
    }


def build_ledger(
    records: list[ParcelRecord],
    output_path: str | Path,
    column_mapping_path: str | Path | None = None,
    site_name: str = "",
) -> Path:
    columns = load_column_mapping(column_mapping_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "토지조서"

    if site_name:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        title_cell = ws.cell(row=1, column=1, value=f"토지조서 - {site_name}")
        title_cell.font = Font(bold=True, size=13)
        title_cell.alignment = Alignment(horizontal="center")
        header_row = 2
    else:
        header_row = 1

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    fail_fill = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")

    area_col_idx = next(i for i, (field, _) in enumerate(columns, start=1) if field == "면적_m2")
    py_col_idx = next(i for i, (field, _) in enumerate(columns, start=1) if field == "면적_평")

    row_idx = header_row
    area_sum = 0.0
    py_sum = 0.0
    excluded_count = 0

    for record in records:
        row_idx += 1
        values = _row_values(record)
        for col_idx, (field, _) in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values.get(field))
            if record.조회상태 == QueryStatus.FAILED:
                cell.fill = fail_fill

        if record.면적_m2 is not None:
            area_sum += record.면적_m2
            py_sum += record.면적_평 or (record.면적_m2 / 3.305785)
        else:
            excluded_count += 1

    # PRD FR-6: 합계 행 (합계에서 빠진 필지 수 명시)
    row_idx += 1
    total_label = f"합계 (조회 실패 {excluded_count}필지 제외)" if excluded_count else "합계"
    ws.cell(row=row_idx, column=1, value=total_label).font = Font(bold=True)
    ws.cell(row=row_idx, column=area_col_idx, value=round(area_sum, 2)).font = Font(bold=True)
    ws.cell(row=row_idx, column=py_col_idx, value=round(py_sum, 2)).font = Font(bold=True)

    for col_idx, (_, header) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(header) * 2)

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
